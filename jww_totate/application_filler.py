import os
import datetime
import re
import time
import openpyxl

def export_selected_sheets_to_pdf(excel_path, sheet_names, pdf_path):
    """ 指定したExcelファイルの指定されたシート群をPDFとして出力する """
    try:
        import win32com.client
        import pythoncom
        pythoncom.CoInitialize()

        excel = win32com.client.DispatchEx("Excel.Application")
        excel.Visible = False
        excel.DisplayAlerts = False
        
        abs_excel_path = os.path.abspath(excel_path)
        abs_pdf_path = os.path.abspath(pdf_path)

        if not os.path.exists(abs_excel_path):
            print(f"⚠️ エラー: PDF化の元となるExcelファイルが見つかりません: {abs_excel_path}")
            return

        wb = excel.Workbooks.Open(abs_excel_path, UpdateLinks=False, ReadOnly=True)
        
        try:
            target_sheets = sheet_names if sheet_names is not None else [sheet.Name for sheet in wb.Worksheets]
            valid_sheets = []
            
            try: excel.PrintCommunication = False
            except: pass
            
            for s in target_sheets:
                try:
                    ws = wb.Worksheets(s)
                    # ★変更: 縦も横も強制的に1ページに収める設定
                    ws.PageSetup.Zoom = False
                    ws.PageSetup.FitToPagesWide = 1
                    ws.PageSetup.FitToPagesTall = 1 # ← Falseから1に変更しました！
                    valid_sheets.append(s)
                except Exception:
                    pass
                    
            try: excel.PrintCommunication = True
            except: pass
            
            if valid_sheets:
                wb.Worksheets(valid_sheets).Select()
                excel.ActiveSheet.ExportAsFixedFormat(0, abs_pdf_path)
                print(f"   => 📄 PDF出力完了: {os.path.basename(pdf_path)}")
            else:
                print(f"⚠️ 指定されたシートが見つからないためPDF化をスキップ: {os.path.basename(excel_path)}")
            
        except Exception as e:
            print(f"   => ⚠️ {os.path.basename(excel_path)} のPDF化処理中にエラーが発生しました: {e}")
        finally:
            wb.Close(False)
            excel.Quit()
    except Exception as e:
        print(f"⚠️ Excel経由でのPDF出力に失敗しました（Windows環境＋Excel必須）: {e}")


def extract_date_parts(date_str):
    """ '2026/08/10' のような文字列から、年・月・日を分割して返す """
    if not date_str:
        return "", "", ""
    parts = re.split(r'[/.-]', date_str.replace(" ", ""))
    if len(parts) >= 3:
        try:
            y = str(int(parts[0]))
            m = str(int(parts[1]))
            d = str(int(parts[2]))
            return y, m, d
        except:
            pass
    return date_str, "", ""


def fill_application_forms(house_name, house_address, common_data, calc_results, envelope_calc_excel_path, house_area, start_date, output_dir="."):
    """
    計算結果や物件情報を申請書のExcelに自動転記し、個別のPDFファイルとして出力する
    """
    print("\n📄 申請書類一式のExcel作成とPDF化を開始します...")
    
    today = datetime.datetime.now()
    year, month, day = str(today.year), str(today.month), str(today.day)

    s_year, s_month, s_day = extract_date_parts(start_date)
    
    region_val = common_data.get("region", "")
    region_num = re.search(r'\d+', region_val.replace("５", "5").replace("６", "6"))
    region_num = region_num.group() if region_num else ""
    
    safe_house_name = "".join(c for c in house_name if c not in r'\/:*?"<>|')

    # ---------------------------------------------------------
    # 1. 01申請書.xlsx の処理
    # ---------------------------------------------------------
    file1 = "01申請書.xlsx"
    out_name1 = os.path.abspath(os.path.join(output_dir, f"01申請書_{safe_house_name}.xlsx"))
    pdf_name1 = os.path.abspath(os.path.join(output_dir, f"01申請書_{safe_house_name}.pdf"))
    
    if os.path.exists(file1):
        try:
            wb = openpyxl.load_workbook(file1)
            if "第一面" in wb.sheetnames:
                ws = wb["第一面"]
                ws['W8'], ws['AA8'], ws['AC8'] = year, month, day
            if "第三面" in wb.sheetnames:
                ws = wb["第三面"]
                ws['J5'] = house_name
                ws['J8'] = house_address
                ws['K9'] = region_num
                ws['J12'] = house_area
                ws['T16'] = s_year
                ws['X16'] = s_month
                ws['AA16'] = s_day
                
            wb.save(out_name1)
            wb.close()
            time.sleep(1)
            
            print(f"✔️ {os.path.basename(out_name1)} を作成しました。")
            export_selected_sheets_to_pdf(out_name1, ["第一面", "第二面", "第三面", "第四面"], pdf_name1)
        except Exception as e:
            print(f"⚠️ {file1} の処理中にエラー: {e}")

    # ---------------------------------------------------------
    # 2. 05委任状.xlsx の処理
    # ---------------------------------------------------------
    file2 = "05委任状.xlsx"
    out_name2 = os.path.abspath(os.path.join(output_dir, f"05委任状_{safe_house_name}.xlsx"))
    pdf_name2 = os.path.abspath(os.path.join(output_dir, f"05委任状_{safe_house_name}.pdf"))
    
    if os.path.exists(file2):
        try:
            wb = openpyxl.load_workbook(file2)
            sheet_target = "委任状" if "委任状" in wb.sheetnames else wb.sheetnames[0]
            ws = wb[sheet_target]
            
            ws['U5'], ws['X5'], ws['Z5'] = year, month, day
            ws['F23'] = house_name
            ws['F27'] = house_address
            
            wb.save(out_name2)
            wb.close()
            time.sleep(1)
            
            print(f"✔️ {os.path.basename(out_name2)} を作成しました。")
            export_selected_sheets_to_pdf(out_name2, [sheet_target], pdf_name2)
        except Exception as e:
            print(f"⚠️ {file2} の処理中にエラー: {e}")

    # ---------------------------------------------------------
    # 3. 06掲載承諾書.xlsx の処理
    # ---------------------------------------------------------
    file3 = "06掲載承諾書.xlsx"
    out_name3 = os.path.abspath(os.path.join(output_dir, f"06掲載承諾書_{safe_house_name}.xlsx"))
    pdf_name3 = os.path.abspath(os.path.join(output_dir, f"06掲載承諾書_{safe_house_name}.pdf"))
    
    if os.path.exists(file3):
        try:
            wb = openpyxl.load_workbook(file3)
            sheet_target = "承諾書" if "承諾書" in wb.sheetnames else wb.sheetnames[0]
            ws = wb[sheet_target]

            ws['W5'], ws['AA5'], ws['AC5'] = year, month, day
            ws['K14'] = house_name
            
            wb.save(out_name3)
            wb.close()
            time.sleep(1)
            
            print(f"✔️ {os.path.basename(out_name3)} を作成しました。")
            export_selected_sheets_to_pdf(out_name3, [sheet_target], pdf_name3)
        except Exception as e:
            print(f"⚠️ {file3} の処理中にエラー: {e}")

    # ---------------------------------------------------------
    # 4. 07設計内説明書.xlsx の処理
    # ---------------------------------------------------------
    file4 = "07設計内説明書20260401.xlsx"
    out_name4 = os.path.abspath(os.path.join(output_dir, f"07設計内容説明書_{safe_house_name}.xlsx"))
    pdf_name4 = os.path.abspath(os.path.join(output_dir, f"07設計内容説明書_{safe_house_name}.pdf"))
    
    if os.path.exists(file4):
        try:
            wb = openpyxl.load_workbook(file4)
            if "第一面" in wb.sheetnames:
                ws = wb["第一面"]
                ws['E5'] = house_name
            
            sheet_target2 = "第二面（住宅）"
            if sheet_target2 in wb.sheetnames:
                ws = wb[sheet_target2]
                ws['I10'] = common_data.get("avg_heat_transfer", "")
                ws['I12'] = common_data.get("cooling_solar_heat_gain", "")
                
                if calc_results:
                    ws['T27'] = float(calc_results.get("design_energy", 0)) if calc_results.get("design_energy") else ""
                    ws['T28'] = float(calc_results.get("standard_energy", 0)) if calc_results.get("standard_energy") else ""
                    ws['T29'] = float(calc_results.get("bei", 0)) if calc_results.get("bei") else ""
            
            wb.save(out_name4)
            wb.close()
            time.sleep(1)
            
            print(f"✔️ {os.path.basename(out_name4)} を作成しました。")
            export_selected_sheets_to_pdf(out_name4, ["第一面", "第二面（住宅）"], pdf_name4)
        except Exception as e:
            print(f"⚠️ {file4} の処理中にエラー: {e}")

    # ---------------------------------------------------------
    # 5. 外皮計算書 のPDF処理
    # ---------------------------------------------------------
    if envelope_calc_excel_path and os.path.exists(envelope_calc_excel_path):
        print(f"✔️ 外皮計算書のPDF化を開始します。")
        time.sleep(1)
        pdf_name_env = os.path.abspath(os.path.join(output_dir, f"外皮計算書_{safe_house_name}.pdf"))
        
        env_sheets_to_export = []
        try:
            wb_env = openpyxl.load_workbook(envelope_calc_excel_path, data_only=True)
            required_sheets = ["共通条件・結果", "Ｂ（屋根・床等）", "Ｃ（基礎）"]
            for s in required_sheets:
                if s in wb_env.sheetnames:
                    env_sheets_to_export.append(s)
            
            direction_sheets = ['Ａ（北）', 'Ａ（北東）', 'Ａ（東）', 'Ａ（南東）', 'Ａ（南）', 'Ａ（南西）', 'Ａ（西）', 'Ａ（北西）']
            for s in direction_sheets:
                if s in wb_env.sheetnames:
                    ws = wb_env[s]
                    has_input = False
                    
                    wall_area = ws['L35'].value
                    if wall_area not in [None, 0, "", "0", "0.0"]:
                        has_input = True
                    elif ws['B8'].value not in [None, ""]:
                        has_input = True
                    
                    if has_input:
                        env_sheets_to_export.append(s)
                        
            wb_env.close()
        except Exception as e:
            print(f"⚠️ 外皮計算書のシート判定に失敗したため全シートを出力します: {e}")
            env_sheets_to_export = None

        export_selected_sheets_to_pdf(envelope_calc_excel_path, env_sheets_to_export, pdf_name_env)